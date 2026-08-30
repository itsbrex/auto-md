from __future__ import annotations

import hashlib
import json
import logging
import mimetypes
import re
import tempfile
import threading
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from markdownify import markdownify

from .config import LIMITS, SUPPORTED_IMAGE_EXTENSIONS
from .models import Document, OcrMode, ProcessingOptions, StagedFile

logger = logging.getLogger(__name__)


class UnsupportedFormat(ValueError):
    pass


TEXT_EXTENSIONS = {
    ".txt",
    ".text",
    ".log",
    ".md",
    ".markdown",
    ".mdown",
    ".rst",
    ".py",
    ".pyw",
    ".js",
    ".jsx",
    ".mjs",
    ".cjs",
    ".ts",
    ".tsx",
    ".java",
    ".c",
    ".h",
    ".cpp",
    ".hpp",
    ".cc",
    ".cs",
    ".go",
    ".rs",
    ".rb",
    ".php",
    ".swift",
    ".kt",
    ".kts",
    ".scala",
    ".dart",
    ".lua",
    ".pl",
    ".pm",
    ".r",
    ".m",
    ".mm",
    ".ex",
    ".exs",
    ".erl",
    ".hrl",
    ".hs",
    ".lhs",
    ".ml",
    ".mli",
    ".clj",
    ".cljs",
    ".vim",
    ".json",
    ".jsonl",
    ".json5",
    ".yaml",
    ".yml",
    ".toml",
    ".xml",
    ".xsl",
    ".xslt",
    ".svg",
    ".csv",
    ".tsv",
    ".sql",
    ".tex",
    ".bib",
    ".ini",
    ".cfg",
    ".conf",
    ".config",
    ".editorconfig",
    ".env",
    ".sh",
    ".bash",
    ".zsh",
    ".fish",
    ".bat",
    ".cmd",
    ".ps1",
    ".css",
    ".scss",
    ".sass",
    ".less",
    ".dockerfile",
    ".containerfile",
    ".gitignore",
    ".gitattributes",
    ".gitmodules",
}

HTML_EXTENSIONS = {".html", ".htm", ".xhtml", ".shtml"}
NOTEBOOK_EXTENSIONS = {".ipynb"}

LANGUAGES = {
    ".py": "python",
    ".pyw": "python",
    ".js": "javascript",
    ".jsx": "jsx",
    ".mjs": "javascript",
    ".cjs": "javascript",
    ".ts": "typescript",
    ".tsx": "tsx",
    ".rs": "rust",
    ".go": "go",
    ".java": "java",
    ".c": "c",
    ".h": "c",
    ".cpp": "cpp",
    ".hpp": "cpp",
    ".cc": "cpp",
    ".cs": "csharp",
    ".rb": "ruby",
    ".php": "php",
    ".swift": "swift",
    ".kt": "kotlin",
    ".kts": "kotlin",
    ".scala": "scala",
    ".dart": "dart",
    ".lua": "lua",
    ".r": "r",
    ".sh": "bash",
    ".bash": "bash",
    ".zsh": "zsh",
    ".fish": "fish",
    ".ps1": "powershell",
    ".bat": "batch",
    ".cmd": "batch",
    ".json": "json",
    ".jsonl": "jsonl",
    ".json5": "json5",
    ".yaml": "yaml",
    ".yml": "yaml",
    ".toml": "toml",
    ".xml": "xml",
    ".svg": "xml",
    ".csv": "csv",
    ".tsv": "tsv",
    ".sql": "sql",
    ".md": "markdown",
    ".markdown": "markdown",
    ".rst": "rst",
    ".css": "css",
    ".scss": "scss",
    ".sass": "sass",
    ".less": "less",
    ".html": "html",
    ".htm": "html",
}

SPECIAL_TEXT_NAMES = {
    "dockerfile",
    "containerfile",
    "makefile",
    "procfile",
    "gemfile",
    "rakefile",
    "license",
    "readme",
    "changelog",
    "contributing",
    "security",
    "code_of_conduct",
}


class Extractor:
    def __init__(self) -> None:
        self._ocr_engine: Any | None = None
        self._ocr_lock = threading.Lock()
        self._pdf_lock = threading.Lock()

    def extract(self, staged: StagedFile, options: ProcessingOptions) -> Document:
        path = staged.path
        suffix = path.suffix.lower()
        sha256 = _sha256_file(path)
        if suffix in HTML_EXTENSIONS:
            content = self._html(path)
            return self._document(staged, content, "text/html", sha256, content_kind="markdown")
        if suffix == ".docx":
            content = self._docx(path)
            return self._document(
                staged,
                content,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                sha256,
                content_kind="markdown",
            )
        if suffix == ".pdf":
            content, warnings = self._pdf(path, options.ocr_mode)
            document = self._document(
                staged, content, "application/pdf", sha256, content_kind="markdown"
            )
            document.warnings.extend(warnings)
            return document
        if suffix in SUPPORTED_IMAGE_EXTENSIONS:
            if options.ocr_mode == OcrMode.NEVER:
                raise UnsupportedFormat("Image OCR is disabled")
            content = self._image(path)
            return self._document(
                staged, content, f"image/{suffix.lstrip('.')}", sha256, content_kind="markdown"
            )
        if suffix in NOTEBOOK_EXTENSIONS:
            content = self._notebook(path, options.include_notebook_outputs)
            return self._document(
                staged, content, "application/x-ipynb+json", sha256, content_kind="markdown"
            )
        if self.is_text(path):
            raw = path.read_bytes()
            content, warning = decode_text(raw)
            media_type = mimetypes.guess_type(path.name)[0] or "text/plain"
            document = self._document(
                staged,
                content,
                media_type,
                sha256,
                language=LANGUAGES.get(suffix, "text"),
            )
            if warning:
                document.warnings.append(warning)
            return document
        raise UnsupportedFormat("Unsupported or binary file")

    def is_text(self, path: Path) -> bool:
        suffix = path.suffix.lower()
        if suffix in TEXT_EXTENSIONS:
            return True
        stem = path.name.lower()
        if stem in SPECIAL_TEXT_NAMES or any(
            stem.startswith(f"{name}.") for name in SPECIAL_TEXT_NAMES
        ):
            return True
        sample = path.read_bytes()[:8192]
        if not sample or b"\x00" in sample:
            return False
        control = sum(byte < 9 or 13 < byte < 32 for byte in sample)
        return control / len(sample) < 0.02

    def _document(
        self,
        staged: StagedFile,
        content: str,
        media_type: str,
        sha256: str,
        *,
        content_kind: str = "raw",
        language: str | None = None,
    ) -> Document:
        title = staged.relative_path or staged.source_label
        return Document(
            source_id=staged.source_id,
            source_label=staged.source_label,
            path=staged.relative_path,
            title=title,
            media_type=media_type,
            content=content.strip("\n"),
            content_kind=content_kind,
            language=language,
            source_url=staged.source_url,
            sha256=sha256,
        )

    def _html(self, path: Path) -> str:
        text, _ = decode_text(path.read_bytes())
        soup = BeautifulSoup(text, "html.parser")
        for element in soup(["script", "style", "noscript", "template", "svg", "canvas"]):
            element.decompose()
        content = soup.find("article") or soup.find("main") or soup.body or soup
        for element in content.find_all(["nav", "aside", "footer"]):
            element.decompose()
        return markdownify(str(content), heading_style="ATX", bullets="-").strip()

    def _docx(self, path: Path) -> str:
        from docx import Document as DocxDocument

        document = DocxDocument(path)
        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style else "").lower()
            heading_match = re.match(r"heading\s+(\d+)", style)
            if heading_match:
                level = min(int(heading_match.group(1)), 6)
                parts.append(f"{'#' * level} {text}")
            elif "list" in style:
                marker = "1." if "number" in style else "-"
                parts.append(f"{marker} {text}")
            else:
                parts.append(text)
        for table_index, table in enumerate(document.tables, start=1):
            rows = [[_clean_cell(cell.text) for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            parts.append(f"### Table {table_index}")
            parts.append("| " + " | ".join(rows[0]) + " |")
            parts.append("| " + " | ".join("---" for _ in range(width)) + " |")
            parts.extend("| " + " | ".join(row) + " |" for row in rows[1:])
        return "\n\n".join(parts)

    def _pdf(self, path: Path, ocr_mode: OcrMode) -> tuple[str, list[str]]:
        import pypdfium2 as pdfium  # type: ignore[import-untyped]

        warnings: list[str] = []
        pages: list[str] = []
        with self._pdf_lock:
            pdf = pdfium.PdfDocument(str(path))
            try:
                if len(pdf) > LIMITS.max_pdf_pages:
                    raise ValueError(f"PDF exceeds the {LIMITS.max_pdf_pages}-page limit")
                for index in range(len(pdf)):
                    page = pdf[index]
                    text_page = page.get_textpage()
                    try:
                        text = text_page.get_text_range().strip()
                    finally:
                        text_page.close()
                    needs_ocr = ocr_mode == OcrMode.ALWAYS or (
                        ocr_mode == OcrMode.AUTO and len(re.findall(r"[A-Za-z0-9]", text)) < 20
                    )
                    if needs_ocr:
                        bitmap = page.render(scale=200 / 72)
                        try:
                            image = bitmap.to_pil()
                            with tempfile.NamedTemporaryFile(suffix=".png") as temporary:
                                image.save(temporary.name, "PNG")
                                ocr_text = self._run_ocr(Path(temporary.name))
                            if ocr_text.strip():
                                text = ocr_text
                            elif not text:
                                warnings.append(f"Page {index + 1}: OCR found no text")
                        finally:
                            bitmap.close()
                    pages.append(f"## Page {index + 1}\n\n{text}".rstrip())
                    page.close()
            finally:
                pdf.close()
        return "\n\n".join(pages), warnings

    def _image(self, path: Path) -> str:
        from PIL import Image, ImageOps

        with Image.open(path) as image:
            if image.width * image.height > LIMITS.max_image_pixels:
                raise ValueError("Image exceeds the pixel limit")
            image = ImageOps.exif_transpose(image).convert("RGB")
            with tempfile.NamedTemporaryFile(suffix=".png") as temporary:
                image.save(temporary.name, "PNG")
                text = self._run_ocr(Path(temporary.name))
        if not text.strip():
            raise ValueError("OCR found no text in the image")
        return f"## OCR text\n\n{text.strip()}"

    def _run_ocr(self, path: Path) -> str:
        with self._ocr_lock:
            if self._ocr_engine is None:
                from rapidocr import RapidOCR

                self._ocr_engine = RapidOCR()
            result = self._ocr_engine(str(path))
        if hasattr(result, "txts"):
            values = result.txts or []
            return "\n".join(str(value) for value in values)
        if isinstance(result, tuple):
            result = result[0]
        if isinstance(result, list):
            lines: list[str] = []
            for item in result:
                if isinstance(item, (list, tuple)) and len(item) >= 2:
                    candidate = item[1]
                    if isinstance(candidate, (list, tuple)) and candidate:
                        lines.append(str(candidate[0]))
                    elif isinstance(candidate, str):
                        lines.append(candidate)
            return "\n".join(lines)
        return ""

    def _notebook(self, path: Path, include_outputs: bool) -> str:
        value = json.loads(path.read_text(encoding="utf-8"))
        cells = value.get("cells")
        if not isinstance(cells, list):
            raise ValueError("Notebook has no cells array")
        parts: list[str] = []
        for index, cell in enumerate(cells, start=1):
            if not isinstance(cell, dict):
                continue
            source = _join_notebook_text(cell.get("source", ""))
            cell_type = cell.get("cell_type")
            if cell_type == "markdown":
                parts.append(f"## Cell {index} · Markdown\n\n{source.strip()}")
            elif cell_type == "code":
                fence = markdown_fence(source)
                section = f"## Cell {index} · Code\n\n{fence}python\n{source.rstrip()}\n{fence}"
                if include_outputs:
                    output_text = _notebook_outputs(cell.get("outputs", []))
                    if output_text:
                        output_fence = markdown_fence(output_text)
                        section += (
                            f"\n\n### Output\n\n{output_fence}text\n{output_text}\n{output_fence}"
                        )
                parts.append(section)
        return "\n\n".join(parts)


def decode_text(raw: bytes) -> tuple[str, str | None]:
    if raw.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        return raw.decode("utf-32"), None
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16"), None
    try:
        return raw.decode("utf-8-sig"), None
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace"), "Decoded as Windows-1252 after UTF-8 failed"


def markdown_fence(text: str) -> str:
    longest = max((len(match.group(0)) for match in re.finditer(r"`+", text)), default=0)
    return "`" * max(3, longest + 1)


def _join_notebook_text(value: Any) -> str:
    if isinstance(value, list):
        return "".join(str(item) for item in value)
    return str(value)


def _notebook_outputs(value: Any) -> str:
    if not isinstance(value, list):
        return ""
    parts: list[str] = []
    for output in value:
        if not isinstance(output, dict):
            continue
        if "text" in output:
            parts.append(_join_notebook_text(output["text"]).strip())
        data = output.get("data")
        if isinstance(data, dict) and "text/plain" in data:
            parts.append(_join_notebook_text(data["text/plain"]).strip())
    return "\n\n".join(part for part in parts if part)


def _clean_cell(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().replace("|", "\\|")


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()
